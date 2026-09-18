import pytest
from io import BytesIO
from docx import Document
from app.services.docx_import.parser import parse_docx_timetable

def create_department_docx(main_header, main_rows, legend_header, legend_rows):
    doc = Document()

    # Main Timetable (needs an extra row at the top because parser expects header at index 1)
    table1 = doc.add_table(rows=len(main_rows) + 2, cols=len(main_header))

    # Add dummy title row
    for col_idx in range(len(main_header)):
        table1.cell(0, col_idx).text = "Dummy Title Row"

    # Add header (index 1)
    for col_idx, text in enumerate(main_header):
        table1.cell(1, col_idx).text = text

    # Add rows (index 2+)
    for row_idx, row_data in enumerate(main_rows):
        for col_idx, text in enumerate(row_data):
            table1.cell(row_idx + 2, col_idx).text = text

    # Legend Table
    table2 = doc.add_table(rows=len(legend_rows) + 1, cols=len(legend_header))

    # Add legend header
    for col_idx, text in enumerate(legend_header):
        table2.cell(0, col_idx).text = text

    # Add legend rows
    for row_idx, row_data in enumerate(legend_rows):
        for col_idx, text in enumerate(row_data):
            table2.cell(row_idx + 1, col_idx).text = text

    # Save to BytesIO
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def test_department_agnostic_mca():
    """Test MCA-style timetable with 'Initial' and 'Classroom/ Laboratory' headers."""
    main_header = ["Day", "Section", "08:55 - 09:50", "09:50 - 10:45"]
    main_rows = [
        ["Monday", "I-A", "DBMS\n(VR)\nCA1", "LAB\n(SU)\n(LAB1A)"]
    ]

    legend_header = ["Sl. No.", "Name of Faculty", "Initial", "", "Classroom/ Laboratory", "Location"]
    legend_rows = [
        ["1", "Veena", "VR", "", "CA1", "Floor 1"],
        ["2", "S Uma", "SU", "", "LAB1A", "Floor 2"]
    ]

    docx_file = create_department_docx(main_header, main_rows, legend_header, legend_rows)

    # Read the file
    preview = parse_docx_timetable(docx_file.getvalue(), "MCA", "2026-2027", "MCA_TEST.docx")

    # Verify legends extracted based on dynamic headers
    assert "VR" in preview.faculty_legend
    assert preview.faculty_legend["VR"] == "Veena"

    assert "CA1" in preview.resource_legend
    assert "LAB1A" in preview.resource_legend

    # Find block
    assert len(preview.occupancy_ready_blocks) >= 2

    blocks = sorted(preview.occupancy_ready_blocks, key=lambda b: b.slots[0])
    block1 = blocks[0]  # S2
    block2 = blocks[1]  # S3

    assert "S2" in block1.slots
    assert block1.resource_candidates[0].code == "CA1"
    assert block1.teacher_candidates[0].acronym == "VR"

    assert "S3" in block2.slots
    assert block2.resource_candidates[0].code == "LAB1A"
    assert block2.teacher_candidates[0].acronym == "SU"

def test_department_agnostic_cs():
    """Test completely different department format relying entirely on legends."""
    main_header = ["Day", "Section", "08:55 - 09:50", "09:50 - 10:45"]
    main_rows = [
        ["Monday", "A", "AI\n(ALICE)\nCSE201", "DS\n(BOB)\n(AILAB)"]
    ]

    # Notice different headers: "Acronym" instead of "Initial", "Room" instead of "Classroom"
    legend_header = ["S.N", "Faculty Name", "Acronym", "Room Type", "Room", "Location"]
    legend_rows = [
        ["1", "Alice CS", "ALICE", "Class", "CSE201", "Block A"],
        ["2", "Bob CS", "BOB", "Lab", "AILAB", "Block B"]
    ]

    docx_file = create_department_docx(main_header, main_rows, legend_header, legend_rows)

    preview = parse_docx_timetable(docx_file.getvalue(), "CS", "2026-2027", "CS_TEST.docx")

    # Legends should dynamically pick up ALICE and CSE201 due to generic keyword matches
    assert "ALICE" in preview.faculty_legend
    assert preview.faculty_legend["ALICE"] == "Alice CS"

    assert "CSE201" in preview.resource_legend
    assert "AILAB" in preview.resource_legend

    blocks = sorted(preview.occupancy_ready_blocks, key=lambda b: b.slots[0])
    block1 = blocks[0]
    block2 = blocks[1]

    # Candidates successfully parsed because they are present in the dynamic legend
    assert block1.resource_candidates[0].code == "CSE201"
    assert block1.teacher_candidates[0].acronym == "ALICE"

    assert block2.resource_candidates[0].code == "AILAB"
    assert block2.teacher_candidates[0].acronym == "BOB"
